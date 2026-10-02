"""Analytics for annual Home Coming custom payment links with installments."""
from __future__ import annotations

import calendar
from datetime import date, datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from utils.excel_dates import iso_ymd_to_dd_mm_yyyy

_ANNUAL_ITEM_TYPES = frozenset({"annual_package"})
_ANNUAL_PLANS = frozenset({"quarter_then_monthly", "down_then_emi"})
_EMI_DAY = 27


def _parse_ymd(s: str) -> Optional[date]:
    t = (s or "").strip()[:10]
    if len(t) < 10 or t[4] != "-" or t[7] != "-":
        return None
    try:
        return datetime.strptime(t, "%Y-%m-%d").date()
    except ValueError:
        return None


def _emi_due_date_iso(start_yyyy_mm_dd: str, month_offset: int, emi_day: int = _EMI_DAY) -> str:
    """Same calendar rule as Iris Annual Abundance EMIs (27th when possible)."""
    if not start_yyyy_mm_dd or len(start_yyyy_mm_dd) < 10:
        return ""
    day = int(emi_day) if emi_day else _EMI_DAY
    if day < 1:
        day = _EMI_DAY
    try:
        y = int(start_yyyy_mm_dd[0:4])
        mo = int(start_yyyy_mm_dd[5:7])
        total_m = mo - 1 + month_offset
        y += total_m // 12
        m2 = total_m % 12 + 1
        dim = calendar.monthrange(y, m2)[1]
        d = min(day, dim)
        return f"{y:04d}-{m2:02d}-{d:02d}"
    except (ValueError, TypeError):
        return ""


def is_annual_installment_payment_link(row: dict) -> bool:
    if not row.get("installments_enabled"):
        return False
    item_type = (row.get("item_type") or "").strip().lower()
    if item_type in _ANNUAL_ITEM_TYPES:
        return True
    plan = (row.get("installment_plan") or "").strip().lower()
    if plan in _ANNUAL_PLANS:
        return True
    blob = f"{row.get('title') or ''} {row.get('item_title') or ''}".lower()
    if "home coming" in blob or "homecoming" in blob:
        return True
    if "annual" in blob and ("package" in blob or "year" in blob):
        return True
    return False


def _schedule_anchor_ymd(row: dict, paid_by_number: Dict[int, dict]) -> str:
    start = (row.get("chosen_start_date") or "").strip()[:10]
    if _parse_ymd(start):
        return start
    if paid_by_number:
        first = paid_by_number.get(min(paid_by_number.keys()))
        if first:
            pa = (first.get("paid_at") or "")[:10]
            if _parse_ymd(pa):
                return pa
    created = (row.get("created_at") or "")[:10]
    if _parse_ymd(created):
        return created
    return ""


def _due_ymd_for_installment(row: dict, inst_num: int, anchor: str) -> str:
    if not anchor:
        return ""
    plan = (row.get("installment_plan") or "equal").strip().lower()
    n = int(inst_num)
    if plan in _ANNUAL_PLANS:
        if n <= 1:
            return anchor
        return _emi_due_date_iso(anchor, n - 2, _EMI_DAY)
    return _emi_due_date_iso(anchor, n - 1, _EMI_DAY)


def _slot_status(due_ymd: str, paid_entry: Optional[dict], link_status: str) -> str:
    if paid_entry:
        return "paid"
    if (link_status or "").lower() == "cancelled":
        return "cancelled"
    d = _parse_ymd(due_ymd)
    if not d:
        return "pending"
    today = datetime.now(timezone.utc).date()
    if d < today:
        return "overdue"
    return "due"


def build_annual_installment_analytics(rows: List[dict]) -> Dict[str, Any]:
    today = datetime.now(timezone.utc).date()
    members: List[dict] = []
    month_buckets: Dict[str, Dict[str, int]] = {}

    def bump_month(ym: str, status: str) -> None:
        if not ym:
            return
        b = month_buckets.setdefault(ym, {"paid": 0, "due": 0, "overdue": 0, "pending": 0, "cancelled": 0})
        key = status if status in b else "pending"
        b[key] = int(b.get(key) or 0) + 1

    for row in rows:
        if not is_annual_installment_payment_link(row):
            continue
        payments = row.get("installment_payments") or []
        paid_by_number: Dict[int, dict] = {}
        for p in payments:
            if not isinstance(p, dict):
                continue
            try:
                num = int(p.get("number") or 0)
            except (TypeError, ValueError):
                continue
            if num > 0:
                paid_by_number[num] = p

        amounts = row.get("installment_amounts") or []
        if not isinstance(amounts, list) or len(amounts) < 2:
            from routes.payment_requests import _installment_amounts_for_row

            amounts = _installment_amounts_for_row(row)

        n_total = int(row.get("num_installments") or len(amounts) or 0)
        anchor = _schedule_anchor_ymd(row, paid_by_number)
        link_st = (row.get("status") or "").strip().lower()
        name = (row.get("payer_name") or row.get("recipient_name") or "").strip()
        email = (row.get("payer_email") or row.get("recipient_email") or "").strip().lower()
        slots: List[dict] = []
        for i in range(1, max(1, n_total) + 1):
            due = _due_ymd_for_installment(row, i, anchor)
            pe = paid_by_number.get(i)
            st = _slot_status(due, pe, link_st)
            amt = amounts[i - 1] if i <= len(amounts) else None
            slot = {
                "number": i,
                "amount": round(float(amt), 2) if amt is not None else None,
                "due_date": due,
                "status": st,
                "paid_at": (pe.get("paid_at") if pe else "") or "",
                "paid_amount": round(float(pe.get("amount") or 0), 2) if pe else None,
            }
            slots.append(slot)
            ym = due[:7] if due and len(due) >= 7 else ""
            bump_month(ym, st)

        paid_count = len(paid_by_number)
        members.append(
            {
                "payment_request_id": row.get("id"),
                "title": row.get("title"),
                "recipient_name": name,
                "recipient_email": email,
                "link_status": link_st,
                "installment_plan": row.get("installment_plan"),
                "total_amount": round(float(row.get("amount") or 0), 2),
                "currency": (row.get("currency") or "").upper(),
                "installments_paid": paid_count,
                "num_installments": n_total,
                "chosen_start_date": (row.get("chosen_start_date") or "")[:10],
                "chosen_end_date": (row.get("chosen_end_date") or "")[:10],
                "item_type": row.get("item_type"),
                "item_title": row.get("item_title"),
                "schedule_anchor": anchor,
                "installment_slots": slots,
            }
        )

    summary = {
        "as_of": today.isoformat(),
        "member_count": len(members),
        "links_fully_paid": sum(1 for m in members if m["link_status"] == "paid"),
        "links_partially_paid": sum(1 for m in members if m["link_status"] == "partially_paid"),
        "links_active_unpaid": sum(1 for m in members if m["link_status"] == "active"),
        "links_cancelled": sum(1 for m in members if m["link_status"] == "cancelled"),
        "installment_slots_paid": sum(
            1 for m in members for s in m["installment_slots"] if s["status"] == "paid"
        ),
        "installment_slots_overdue": sum(
            1 for m in members for s in m["installment_slots"] if s["status"] == "overdue"
        ),
    }
    monthly = [
        {"month": k, **v}
        for k, v in sorted(month_buckets.items())
    ]
    return {"summary": summary, "monthly": monthly, "members": members}


def analytics_rows_to_xlsx_bytes(report: Dict[str, Any]) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="5B21B6", end_color="5B21B6", fill_type="solid")
    summary = report.get("summary") or {}
    ws_sum["A1"] = "Annual custom link — installment analytics"
    ws_sum["A1"].font = Font(bold=True, size=12)
    rows_sum = [
        ("As of (UTC date)", summary.get("as_of")),
        ("Members on annual installment links", summary.get("member_count")),
        ("Links fully paid", summary.get("links_fully_paid")),
        ("Links partially paid", summary.get("links_partially_paid")),
        ("Links active (no payment yet)", summary.get("links_active_unpaid")),
        ("Links cancelled", summary.get("links_cancelled")),
        ("Installment slots paid", summary.get("installment_slots_paid")),
        ("Installment slots overdue", summary.get("installment_slots_overdue")),
    ]
    for i, (lab, val) in enumerate(rows_sum, start=3):
        ws_sum.cell(row=i, column=1, value=lab)
        ws_sum.cell(row=i, column=2, value=val)

    ws_m = wb.create_sheet("By month")
    month_hdr = ["Month (due)", "Paid", "Due (upcoming)", "Overdue", "Pending", "Cancelled"]
    for c, lab in enumerate(month_hdr, 1):
        cell = ws_m.cell(row=1, column=c, value=lab)
        cell.font = hdr_font
        cell.fill = hdr_fill
    for r, row in enumerate(report.get("monthly") or [], start=2):
        ws_m.cell(row=r, column=1, value=row.get("month"))
        ws_m.cell(row=r, column=2, value=row.get("paid", 0))
        ws_m.cell(row=r, column=3, value=row.get("due", 0))
        ws_m.cell(row=r, column=4, value=row.get("overdue", 0))
        ws_m.cell(row=r, column=5, value=row.get("pending", 0))
        ws_m.cell(row=r, column=6, value=row.get("cancelled", 0))

    ws = wb.create_sheet("Members")
    max_inst = 12
    headers = [
        "Name",
        "Email",
        "Payment title",
        "Link status",
        "Paid / Total inst.",
        "Total amount",
        "Currency",
        "Annual start",
        "Annual end",
        "Schedule anchor",
        "Plan",
    ]
    for i in range(1, max_inst + 1):
        headers.extend([f"Inst {i} due", f"Inst {i} status", f"Inst {i} paid date"])
    for c, lab in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=lab)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for r_idx, m in enumerate(report.get("members") or [], start=2):
        ws.cell(row=r_idx, column=1, value=m.get("recipient_name"))
        ws.cell(row=r_idx, column=2, value=m.get("recipient_email"))
        ws.cell(row=r_idx, column=3, value=m.get("title"))
        ws.cell(row=r_idx, column=4, value=m.get("link_status"))
        ws.cell(
            row=r_idx,
            column=5,
            value=f"{m.get('installments_paid')}/{m.get('num_installments')}",
        )
        ws.cell(row=r_idx, column=6, value=m.get("total_amount"))
        ws.cell(row=r_idx, column=7, value=m.get("currency"))
        ws.cell(row=r_idx, column=8, value=iso_ymd_to_dd_mm_yyyy(m.get("chosen_start_date")))
        ws.cell(row=r_idx, column=9, value=iso_ymd_to_dd_mm_yyyy(m.get("chosen_end_date")))
        ws.cell(row=r_idx, column=10, value=iso_ymd_to_dd_mm_yyyy(m.get("schedule_anchor")))
        ws.cell(row=r_idx, column=11, value=m.get("installment_plan"))
        slots = {int(s["number"]): s for s in (m.get("installment_slots") or []) if s.get("number")}
        col = 12
        for i in range(1, max_inst + 1):
            s = slots.get(i) or {}
            ws.cell(row=r_idx, column=col, value=iso_ymd_to_dd_mm_yyyy(s.get("due_date")))
            ws.cell(row=r_idx, column=col + 1, value=s.get("status"))
            ws.cell(row=r_idx, column=col + 2, value=iso_ymd_to_dd_mm_yyyy((s.get("paid_at") or "")[:10]))
            col += 3

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
